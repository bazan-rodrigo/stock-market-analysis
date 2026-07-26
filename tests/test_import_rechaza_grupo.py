"""El import de señales/estrategias RECHAZA la funcionalidad de grupo removida
(señales source=group, componentes con Alcance de grupo): perderla en silencio
—importándola como señal de activo o descartando el scope— cambiaría el
resultado. Y el import normal (sin grupo) sigue funcionando."""
import json
from io import BytesIO

import openpyxl
import pytest
import sqlalchemy as sa

from app.database import Base, Session, engine, get_session
from app.models import signal_store


@pytest.fixture()
def db():
    import app.models  # noqa: F401
    Session.remove()
    Base.metadata.create_all(engine)
    tables = ("strategy_component", "strategy", "`signal`",
              "signal_eval_log", "indicator_definitions")
    with engine.begin() as conn:
        for t in tables:
            conn.execute(sa.text(f"DELETE FROM {t}"))
    yield
    sig, strat = signal_store._list_dynamic_tables()
    with engine.begin() as conn:
        for name in list(sig.values()) + list(strat.values()):
            conn.execute(sa.text(f"DROP TABLE IF EXISTS {name}"))
        for t in tables:
            conn.execute(sa.text(f"DELETE FROM {t}"))
    for name in list(sig.values()) + list(strat.values()):
        if name in signal_store._meta.tables:
            signal_store._meta.remove(signal_store._meta.tables[name])
    Session.remove()


def _xlsx(sheets):
    """sheets: [(title, [fila, ...]), ...] — la primera es la hoja activa."""
    wb = openpyxl.Workbook()
    for i, (title, rows) in enumerate(sheets):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = title
        for r in rows:
            ws.append(r)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _add_indicator():
    from app.models.indicator_definition import IndicatorDefinition
    s = get_session()
    s.add(IndicatorDefinition(code="trend_daily", name="t", category="t",
                              type="str", keep_history=True))
    s.commit()


_SIG_HEADERS = ["key", "name", "source", "group_type", "indicator_key",
                "formula_type", "params", "publica"]
_MAP = json.dumps({"map": {"bullish": 80}})


def test_import_senal_source_group_se_rechaza(db):
    from app.services import signal_service
    _add_indicator()

    data = _xlsx([("Señales", [
        _SIG_HEADERS,
        ["g_sig", "Grupo", "group", "sector", "regime_score_d",
         "range", json.dumps({"min": -100, "max": 100}), "no"],
    ])])
    res = signal_service.import_signals_excel(data)
    assert res[0]["status"] == "error"
    assert "grupo" in res[0]["detail"].lower()

    # nada quedó escrito
    from app.models import SignalDefinition
    assert get_session().query(SignalDefinition).count() == 0


def test_import_senal_de_activo_sigue_funcionando(db):
    from app.services import signal_service
    from app.models import SignalDefinition
    _add_indicator()

    data = _xlsx([("Señales", [
        _SIG_HEADERS,
        ["a_sig", "Activo", "asset", "", "trend_daily", "discrete_map", _MAP, "no"],
    ])])
    res = signal_service.import_signals_excel(data)
    assert res[0]["status"] == "ok"
    sig = get_session().query(SignalDefinition).one()
    assert sig.key == "a_sig"


def test_import_estrategia_con_alcance_de_grupo_se_rechaza(db):
    from app.services import signal_service, strategy_service
    _add_indicator()
    signal_service.save_signal(
        key="s_ok", name="S", formula_type="discrete_map",
        params_json=_MAP, indicator_key="trend_daily", is_public=True)

    data = _xlsx([
        ("Estrategias", [["name", "description", "filter_conditions", "publica"],
                         ["E1", "", "", "no"]]),
        ("Componentes", [
            ["strategy_name", "signal_key", "weight", "scope", "group_type", "group_id"],
            ["E1", "s_ok", 1, "own_group", "sector", ""],
        ]),
    ])
    res = strategy_service.import_strategies_excel(data)
    assert res[0]["status"] == "error"
    assert "alcance" in res[0]["detail"].lower()

    from app.models import Strategy
    assert get_session().query(Strategy).count() == 0


def test_import_estrategia_de_activo_sigue_funcionando(db):
    from app.services import signal_service, strategy_service
    from app.models import Strategy
    _add_indicator()
    signal_service.save_signal(
        key="s_ok", name="S", formula_type="discrete_map",
        params_json=_MAP, indicator_key="trend_daily", is_public=True)

    data = _xlsx([
        ("Estrategias", [["name", "description", "filter_conditions", "publica"],
                         ["E1", "", "", "no"]]),
        ("Componentes", [
            ["strategy_name", "signal_key", "weight"],
            ["E1", "s_ok", 2],
        ]),
    ])
    res = strategy_service.import_strategies_excel(data)
    assert res[0]["status"] == "ok"
    strat = get_session().query(Strategy).one()
    assert strat.name == "E1"
    assert len(strat.components) == 1 and strat.components[0].weight == 2


def test_import_planilla_de_senales_en_estrategias_se_rechaza(db):
    """Subir la planilla de SEÑALES en la pantalla de estrategias (comparten
    name/description) creaba estrategias vacías con los nombres de las señales,
    sin avisar. Ahora se detecta por las columnas propias de señales y se rechaza."""
    from app.services import strategy_service
    from app.models import Strategy

    data = _xlsx([   # headers de la planilla de SEÑALES, no de estrategias
        ("Señales", [
            ["key", "name", "description", "indicator_key", "formula_type",
             "params", "publica"],
            ["rsi_sig", "RSI señal", "desc", "trend_daily", "discrete_map",
             _MAP, "si"],
        ]),
    ])
    res = strategy_service.import_strategies_excel(data)
    assert res[0]["status"] == "error"
    assert "señales" in res[0]["detail"].lower()
    assert get_session().query(Strategy).count() == 0


def test_import_estrategia_sin_componentes_se_rechaza(db):
    """Una estrategia sin componentes (planilla sin hoja 'Componentes', o con
    strategy_name que no matchea) no puntúa nada → el import la rechaza en vez
    de crearla vacía en silencio (antes importaba OK sin avisar)."""
    from app.services import strategy_service
    from app.models import Strategy

    data = _xlsx([   # solo hoja Estrategias, sin Componentes
        ("Estrategias", [["name", "description", "filter_conditions", "publica"],
                         ["E1", "", "", "no"]]),
    ])
    res = strategy_service.import_strategies_excel(data)
    assert res[0]["status"] == "error"
    assert "sin componentes" in res[0]["detail"].lower()
    assert get_session().query(Strategy).count() == 0
