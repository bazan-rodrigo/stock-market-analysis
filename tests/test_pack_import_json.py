"""Import de packs JSON: equivalencia con el camino Excel y resolución de los
atributos por nombre.

Lo que fija este archivo es que los dos formatos terminan en el MISMO estado de
la base. Si divergen, un pack validado en JSON podría fallar al convertirlo a
planilla (o peor: importar algo distinto sin avisar).
"""
import json
from io import BytesIO

import openpyxl
import pytest
import sqlalchemy as sa

from app.database import Base, Session, engine, get_session
from app.models import signal_store
from app.services import pack_service as ps


@pytest.fixture()
def db():
    import app.models  # noqa: F401
    Session.remove()
    Base.metadata.create_all(engine)
    tablas = ("strategy_component", "strategy", "`signal`", "signal_eval_log",
              "indicator_definitions", "sectors")
    with engine.begin() as conn:
        for t in tablas:
            conn.execute(sa.text(f"DELETE FROM {t}"))
    yield
    sig, strat = signal_store._list_dynamic_tables()
    with engine.begin() as conn:
        for name in list(sig.values()) + list(strat.values()):
            conn.execute(sa.text(f"DROP TABLE IF EXISTS {name}"))
        for t in tablas:
            conn.execute(sa.text(f"DELETE FROM {t}"))
    for name in list(sig.values()) + list(strat.values()):
        if name in signal_store._meta.tables:
            signal_store._meta.remove(signal_store._meta.tables[name])
    Session.remove()


def _sembrar_catalogo():
    from app.models import Sector
    from app.models.indicator_definition import IndicatorDefinition
    s = get_session()
    s.add(IndicatorDefinition(code="trend_daily", name="Tendencia",
                              category="tendencia", type="str", keep_history=True))
    s.add(IndicatorDefinition(code="rsi_daily", name="RSI", category="momentum",
                              type="num", keep_history=True))
    s.add(Sector(name="Technology"))
    s.commit()
    return s.query(Sector).filter(Sector.name == "Technology").one().id


PACK = {
    "spec_version": 1,
    "pack": "prueba",
    "signals": [
        {"key": "p_tendencia", "name": "Tendencia", "description": "d",
         "indicator_key": "trend_daily", "formula_type": "discrete_map",
         "params": {"map": {"bullish": 100, "bearish": -100}},
         "publica": True},
        {"key": "p_rsi", "name": "RSI", "indicator_key": "rsi_daily",
         "formula_type": "range", "params": {"min": 70, "max": 30, "clamp": True},
         "publica": True},
    ],
    "strategies": [{
        "name": "P1", "description": "e", "publica": True,
        "filter": {"op": "AND", "children": [
            {"cond": {"left": {"type": "attribute", "key": "sector"},
                      "operator": "in",
                      "right": {"type": "const", "value": ["Technology"]}}}]},
        "components": [{"signal_key": "p_tendencia", "weight": 3},
                       {"signal_key": "p_rsi", "weight": 1}],
    }],
}


def _json_bytes(pack) -> bytes:
    return json.dumps(pack, ensure_ascii=False).encode("utf-8")


def _xlsx_bytes(hojas) -> bytes:
    wb = openpyxl.Workbook()
    for i, (titulo, filas) in enumerate(hojas):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = titulo
        for fila in filas:
            ws.append(fila)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _xlsx_del_pack(pack):
    """Las mismas dos planillas que generaría scripts/pack_from_json.py."""
    filas_sig = [[f.get(c) for c in ps.SIGNAL_COLUMNS]
                 for f in ps.signal_rows_from_pack(pack)]
    rows_s, rows_c = ps.strategy_rows_from_pack(pack)
    senales = _xlsx_bytes([("Señales", [list(ps.SIGNAL_COLUMNS)] + filas_sig)])
    estrategias = _xlsx_bytes([
        ("Estrategias", [list(ps.STRATEGY_COLUMNS)]
         + [[f.get(c) for c in ps.STRATEGY_COLUMNS] for f in rows_s]),
        ("Componentes", [list(ps.COMPONENT_COLUMNS)]
         + [[f.get(c) for c in ps.COMPONENT_COLUMNS] for f in rows_c]),
    ])
    return senales, estrategias


def _estado():
    """Lo escrito, en forma comparable entre corridas."""
    from app.models import SignalDefinition, Strategy
    s = get_session()
    señales = sorted(
        (sg.key, sg.name, sg.description, sg.indicator_key, sg.formula_type,
         json.dumps(json.loads(sg.params), sort_keys=True), sg.is_public)
        for sg in s.query(SignalDefinition).all())
    estrategias = sorted(
        (st.name, st.description, json.dumps(json.loads(st.filter_conditions),
                                             sort_keys=True), st.is_public,
         tuple(sorted((c.signal_id, c.weight) for c in st.components)))
        for st in s.query(Strategy).all())
    return señales, estrategias


def test_import_json_escribe_lo_esperado(db):
    from app.models import SignalDefinition, Strategy
    from app.services import signal_service, strategy_service

    sector_id = _sembrar_catalogo()

    res = signal_service.import_signals_file(_json_bytes(PACK), "prueba.json", acting_is_admin=True)
    assert all(r["status"] == "ok" for r in res), res
    res = strategy_service.import_strategies_file(_json_bytes(PACK), "prueba.json")
    assert all(r["status"] == "ok" for r in res), res

    s = get_session()
    assert s.query(SignalDefinition).count() == 2
    strat = s.query(Strategy).one()
    assert len(strat.components) == 2

    # El nombre del sector quedó resuelto al id de ESTA instalación
    arbol = json.loads(strat.filter_conditions)
    assert arbol["children"][0]["cond"]["right"]["value"] == [sector_id]


def test_json_y_excel_dejan_la_base_identica(db):
    """La equivalencia de los dos caminos: mismo pack, mismo resultado."""
    from app.services import signal_service, strategy_service

    _sembrar_catalogo()
    signal_service.import_signals_file(_json_bytes(PACK), "p.json", acting_is_admin=True)
    strategy_service.import_strategies_file(_json_bytes(PACK), "p.json")
    desde_json = _estado()

    # Borrar y repetir por el camino Excel. Los componentes se borran a mano:
    # el delete masivo del ORM no dispara el cascade y sqlite no aplica el FK,
    # así que quedarían huérfanos y se los quedaría la estrategia nueva (que
    # reusa el id recién liberado).
    from app.models import SignalDefinition, Strategy, StrategyComponent
    s = get_session()
    s.query(StrategyComponent).delete()
    s.query(Strategy).delete()
    s.query(SignalDefinition).delete()
    s.commit()

    senales_xlsx, estrategias_xlsx = _xlsx_del_pack(PACK)
    signal_service.import_signals_file(senales_xlsx, "senales.xlsx", acting_is_admin=True)
    strategy_service.import_strategies_file(estrategias_xlsx, "estrategia.xlsx")
    desde_excel = _estado()

    # Los ids de señal cambian entre corridas: se comparan los componentes por
    # posición dentro de cada estrategia, no por id absoluto.
    assert desde_json[0] == desde_excel[0]
    assert [e[:4] for e in desde_json[1]] == [e[:4] for e in desde_excel[1]]
    assert [sorted(w for _sid, w in e[4]) for e in desde_json[1]] \
        == [sorted(w for _sid, w in e[4]) for e in desde_excel[1]]


def test_sector_inexistente_rechaza_el_pack_entero(db):
    """Un nombre de catálogo que no existe en la instalación de destino frena
    el import: sin esto el filtro quedaría comparando contra un texto que
    ningún activo tiene y la estrategia saldría vacía, sin error."""
    from app.models import Strategy
    from app.services import signal_service, strategy_service

    _sembrar_catalogo()
    signal_service.import_signals_file(_json_bytes(PACK), "p.json", acting_is_admin=True)

    malo = json.loads(json.dumps(PACK))
    malo["strategies"][0]["filter"]["children"][0]["cond"]["right"]["value"] = ["Minería"]
    res = strategy_service.import_strategies_file(_json_bytes(malo), "p.json")

    assert res[0]["status"] == "error"
    assert "no existe" in res[0]["detail"]
    assert get_session().query(Strategy).count() == 0


def test_ids_de_sector_siguen_aceptandose(db):
    """Lo que exporta la app trae ids: reimportarlo tiene que seguir andando."""
    from app.models import Strategy
    from app.services import signal_service, strategy_service

    sector_id = _sembrar_catalogo()
    signal_service.import_signals_file(_json_bytes(PACK), "p.json", acting_is_admin=True)

    con_id = json.loads(json.dumps(PACK))
    con_id["strategies"][0]["filter"]["children"][0]["cond"]["right"]["value"] = [sector_id]
    res = strategy_service.import_strategies_file(_json_bytes(con_id), "p.json")

    assert res[0]["status"] == "ok", res
    arbol = json.loads(get_session().query(Strategy).one().filter_conditions)
    assert arbol["children"][0]["cond"]["right"]["value"] == [sector_id]


def test_pack_de_solo_estrategias_en_la_pantalla_de_señales_avisa(db):
    from app.services import signal_service
    solo_estrategias = {"strategies": PACK["strategies"]}
    with pytest.raises(ps.PackError, match="Estrategias"):
        signal_service.import_signals_file(_json_bytes(solo_estrategias), "p.json", acting_is_admin=True)


def test_señal_sin_indicador_se_rechaza(db):
    """Una señal sin indicator_key nunca puntúa, y una señal sin valor no
    cuenta en el promedio (no cuenta como cero): se rechaza en vez de crearla
    muda."""
    from app.models import SignalDefinition
    from app.services import signal_service

    _sembrar_catalogo()
    malo = json.loads(json.dumps(PACK))
    malo["signals"][0]["indicator_key"] = ""
    res = signal_service.import_signals_file(_json_bytes(malo), "p.json", acting_is_admin=True)

    assert any(r["status"] == "error" and "indicator_key" in r["detail"]
               for r in res), res
    assert get_session().query(SignalDefinition).count() == 0


def test_catalogo_exportado_sirve_para_validar_offline(db):
    """El círculo completo del estándar: lo que baja el botón «Catálogo» es lo
    que consume `scripts/validate_pack.py`. Si el catálogo cambiara de forma,
    el validador quedaría dando OK sin verificar nada."""
    from app.services import signal_service

    _sembrar_catalogo()
    signal_service.import_signals_file(_json_bytes(PACK), "p.json", acting_is_admin=True)

    catalogo = json.loads(ps.catalog_bytes().decode("utf-8"))

    codigos = {i["code"]: i for i in catalogo["indicators"]}
    assert codigos["trend_daily"]["type"] == "str"
    assert codigos["trend_daily"]["values"], "faltan las categorías posibles"
    assert codigos["last_close"]["virtual"] is True, (
        "los indicadores virtuales se aceptan como indicator_key: sin ellos en "
        "el catálogo, quien escribe un pack no sabe que existen")
    # El hueco encabeza la lista de cada atributo: es un valor más del
    # catálogo, y sin él un pack no podría pedir "que tenga sector"
    assert catalogo["attributes"]["sector"] == ["(sin sector)", "Technology"]
    assert catalogo["attributes"]["synthetic"] == ["(no sintético)"]
    assert {s["key"] for s in catalogo["signals"]} == {"p_tendencia", "p_rsi"}

    # Y con ESE catálogo, el pack valida sin errores ni omisiones
    resultado = ps.validate_pack(PACK, catalogo)
    assert resultado["errors"] == [], resultado["errors"]
    assert resultado["skipped"] == []

    # …y un pack con un indicador inexistente se detecta offline
    malo = json.loads(json.dumps(PACK))
    malo["signals"][0]["indicator_key"] = "no_existe"
    assert any("no existe" in e
               for e in ps.validate_pack(malo, catalogo)["errors"])


def test_reimportar_actualiza_y_no_duplica(db):
    from app.models import SignalDefinition
    from app.services import signal_service

    _sembrar_catalogo()
    signal_service.import_signals_file(_json_bytes(PACK), "p.json", acting_is_admin=True)

    otro = json.loads(json.dumps(PACK))
    otro["signals"][0]["name"] = "Tendencia v2"
    res = signal_service.import_signals_file(_json_bytes(otro), "p.json", acting_is_admin=True)

    assert all(r["status"] == "ok" for r in res)
    s = get_session()
    assert s.query(SignalDefinition).count() == 2
    assert s.query(SignalDefinition).filter(
        SignalDefinition.key == "p_tendencia").one().name == "Tendencia v2"
